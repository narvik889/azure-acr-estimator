#!/usr/bin/env python3
"""
Build the ACR projection workbook (Inputs / Projection / Summary tabs) from an
already-computed projection (the output of calculate_acr.py).

Formulas reference the Inputs tab live -- nothing here is a pasted number.
Verify with scripts/office/soffice.py-based recalc.py (from the xlsx skill)
after generating: zero formula errors is required before handing this off.

Usage:
    python calculate_acr.py --input workloads.json --output acr_projection.json
    python build_workbook.py --input acr_projection.json --output acr_workbook.xlsx
"""
import argparse
import json
from datetime import date

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins

BLUE = Font(name='Arial', color='0000FF')
BLACK = Font(name='Arial', color='000000')
BOLD = Font(name='Arial', bold=True)
NOTE = Font(name='Arial', italic=True, size=9)
HEADER_FILL = PatternFill(start_color='D9D9D9', end_color='D9D9D9', fill_type='solid')
CURRENCY = '$#,##0.00'
PERCENT = '0.0%'


def fit_sheet_to_page(ws, landscape=True):
    """Fix for the truncation/multi-page-split bug found during testing:
    a wide table with default page setup splits across print pages and cuts
    off mid-column when exported to PDF or printed. Scale to fit width."""
    ws.page_setup.orientation = 'landscape' if landscape else 'portrait'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.4, right=0.4, top=0.5, bottom=0.5)
    ws.freeze_panes = 'A2'  # keep header row visible when scrolling


def build_workbook(projection: dict) -> openpyxl.Workbook:
    items = [i for i in projection['items'] if 'error' not in i]
    errored = [i for i in projection['items'] if 'error' in i]
    months = projection['months']

    wb = openpyxl.Workbook()

    # ---------- Inputs tab ----------
    ws_in = wb.active
    ws_in.title = 'Inputs'
    headers = ['Item Name', 'Service', 'Region', 'SKU', 'Unit Price', 'Unit of Measure',
               'Quantity', 'Hours/Month', 'Monthly Growth Rate', 'Monthly Baseline Cost']
    for col, h in enumerate(headers, start=1):
        c = ws_in.cell(row=1, column=col, value=h)
        c.font = BOLD
        c.fill = HEADER_FILL

    for i, item in enumerate(items, start=2):
        ws_in.cell(row=i, column=1, value=item['name']).font = BLACK
        ws_in.cell(row=i, column=2, value=item['service_name']).font = BLACK
        ws_in.cell(row=i, column=3, value=item['region']).font = BLACK
        ws_in.cell(row=i, column=4, value=item.get('matched_sku', '')).font = BLACK
        p = ws_in.cell(row=i, column=5, value=item['unit_price']); p.font = BLUE; p.number_format = CURRENCY
        ws_in.cell(row=i, column=6, value=item.get('unit_of_measure', '')).font = BLACK
        q = ws_in.cell(row=i, column=7, value=item['quantity']); q.font = BLUE
        h = ws_in.cell(row=i, column=8, value=item['hours_per_month']); h.font = BLUE
        g = ws_in.cell(row=i, column=9, value=item['monthly_growth_rate']); g.font = BLUE; g.number_format = PERCENT
        baseline = ws_in.cell(row=i, column=10, value=f'=E{i}*G{i}*H{i}')
        baseline.font = BLACK
        baseline.number_format = CURRENCY

    note_row = len(items) + 3
    ws_in.cell(row=note_row, column=1,
               value=f'Pricing source: Azure Retail Prices API (Consumption/pay-as-you-go), retrieved via MCP tool, {date.today().isoformat()}.').font = NOTE
    ws_in.cell(row=note_row + 1, column=1,
               value='Blue = hardcoded input (live pricing lookup or user-provided assumption). Black = formula.').font = NOTE
    if errored:
        ws_in.cell(row=note_row + 2, column=1,
                    value=f'{len(errored)} item(s) excluded -- price could not be resolved. See calculate_acr.py output for details.').font = Font(name='Arial', italic=True, size=9, color='CC0000')

    # Item Name gets real width so it doesn't truncate in the cell view
    ws_in.column_dimensions['A'].width = 30
    for col in range(2, 11):
        ws_in.column_dimensions[get_column_letter(col)].width = 20
    fit_sheet_to_page(ws_in)

    # ---------- Projection tab ----------
    ws_proj = wb.create_sheet('Projection')
    proj_headers = ['Month'] + [item['name'] for item in items] + ['Total']
    for col, h in enumerate(proj_headers, start=1):
        c = ws_proj.cell(row=1, column=col, value=h)
        c.font = BOLD
        c.fill = HEADER_FILL

    for m in range(1, months + 1):
        row = m + 1
        ws_proj.cell(row=row, column=1, value=m).font = BLACK
        for idx, _ in enumerate(items, start=2):
            input_row = idx
            formula = (f'=Inputs!$E${input_row}*Inputs!$G${input_row}*Inputs!$H${input_row}'
                       f'*(1+Inputs!$I${input_row})^(Projection!$A{row}-1)')
            cell = ws_proj.cell(row=row, column=idx, value=formula)
            cell.font = BLACK
            cell.number_format = CURRENCY
        total_col = len(items) + 2
        first_col_letter = get_column_letter(2)
        last_col_letter = get_column_letter(len(items) + 1)
        total_cell = ws_proj.cell(row=row, column=total_col,
                                   value=f'=SUM({first_col_letter}{row}:{last_col_letter}{row})')
        total_cell.font = BOLD
        total_cell.number_format = CURRENCY

    ws_proj.column_dimensions['A'].width = 10
    for col in range(2, len(items) + 3):
        ws_proj.column_dimensions[get_column_letter(col)].width = 24
    fit_sheet_to_page(ws_proj)

    # ---------- Summary tab ----------
    ws_sum = wb.create_sheet('Summary')
    ws_sum['A1'] = 'Azure Consumed Revenue (ACR) Projection Summary'
    ws_sum['A1'].font = Font(name='Arial', bold=True, size=14)

    total_col_letter = get_column_letter(len(items) + 2)
    ws_sum['A3'] = f'{min(months, 12)}-month ACR'
    ws_sum['A3'].font = BOLD
    end_row_12mo = min(months, 12) + 1
    ws_sum['B3'] = f'=SUM(Projection!${total_col_letter}$2:${total_col_letter}${end_row_12mo})'
    ws_sum['B3'].font = BLACK
    ws_sum['B3'].number_format = CURRENCY

    row_ptr = 4
    if months >= 24:
        ws_sum[f'A{row_ptr}'] = '24-month ACR'
        ws_sum[f'A{row_ptr}'].font = BOLD
        ws_sum[f'B{row_ptr}'] = f'=SUM(Projection!${total_col_letter}$2:${total_col_letter}$25)'
        ws_sum[f'B{row_ptr}'].font = BLACK
        ws_sum[f'B{row_ptr}'].number_format = CURRENCY
        row_ptr += 1

    ws_sum[f'A{row_ptr}'] = 'Month 1 run-rate'
    ws_sum[f'A{row_ptr}'].font = BOLD
    ws_sum[f'B{row_ptr}'] = f'=Projection!${total_col_letter}$2'
    ws_sum[f'B{row_ptr}'].font = BLACK
    ws_sum[f'B{row_ptr}'].number_format = CURRENCY
    row_ptr += 1

    ws_sum[f'A{row_ptr}'] = 'Final month run-rate'
    ws_sum[f'A{row_ptr}'].font = BOLD
    ws_sum[f'B{row_ptr}'] = f'=Projection!${total_col_letter}${months + 1}'
    ws_sum[f'B{row_ptr}'].font = BLACK
    ws_sum[f'B{row_ptr}'].number_format = CURRENCY
    row_ptr += 2

    ws_sum[f'A{row_ptr}'] = 'Notes'
    ws_sum[f'A{row_ptr}'].font = BOLD
    row_ptr += 1
    ws_sum[f'A{row_ptr}'] = f'Pricing source: Azure Retail Prices API (Consumption/pay-as-you-go), retrieved via MCP tool, {date.today().isoformat()}.'
    ws_sum[f'A{row_ptr}'].font = NOTE
    row_ptr += 1
    ws_sum[f'A{row_ptr}'] = 'Reflects retail pay-as-you-go rates, not negotiated/EA pricing, reservations, or savings plans.'
    ws_sum[f'A{row_ptr}'].font = NOTE
    row_ptr += 1
    growth_notes = '; '.join(f"{item['name']} {item['monthly_growth_rate']*100:.1f}%/month" for item in items)
    ws_sum[f'A{row_ptr}'] = f'Growth assumptions: {growth_notes}.'
    ws_sum[f'A{row_ptr}'].font = NOTE
    if errored:
        row_ptr += 1
        ws_sum[f'A{row_ptr}'] = f'{len(errored)} item(s) excluded from this workbook -- price could not be resolved.'
        ws_sum[f'A{row_ptr}'].font = Font(name='Arial', italic=True, size=9, color='CC0000')

    ws_sum.column_dimensions['A'].width = 26
    ws_sum.column_dimensions['B'].width = 18
    fit_sheet_to_page(ws_sum, landscape=False)

    return wb


def main():
    ap = argparse.ArgumentParser(description='Build the ACR projection workbook from calculate_acr.py output')
    ap.add_argument('--input', required=True, help='Path to acr_projection.json (output of calculate_acr.py)')
    ap.add_argument('--output', required=True, help='Path to write the .xlsx workbook')
    args = ap.parse_args()

    with open(args.input) as f:
        projection = json.load(f)

    wb = build_workbook(projection)
    wb.save(args.output)
    print(f'Wrote workbook to {args.output}')
    print('IMPORTANT: run recalc.py on this file next (see xlsx skill) -- openpyxl writes formulas with no cached values.')


if __name__ == '__main__':
    main()
